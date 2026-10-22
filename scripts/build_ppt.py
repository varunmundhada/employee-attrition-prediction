"""Builds the presentation in the college deck style (SIT-N title slide, yellow title bars),
using the earlier CA deck as the template and the numbers from the latest training run.

python -m scripts.build_ppt [template.pptx]
"""

import copy
import json
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from src import config

TEMPLATE = (r"C:\Users\Varun Mundhada\Downloads\One Year Internship\One Year Internship"
            r"\CA6 Internship PPT - Varun Mundhada_1.pptx")
OUT = config.ROOT / "ppt" / "Employee_Attrition_Prediction_Varun_Mundhada.pptx"
FIG = config.FIGURES_DIR
SHOTS = config.ROOT / "docs" / "screenshots"

RED = RGBColor(0xFF, 0x00, 0x00)
DARK = RGBColor(0x1F, 0x1F, 0x1F)
HEAD_FILL = RGBColor(0xFF, 0xC0, 0x00)
ALT_FILL = RGBColor(0xFF, 0xF6, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


# ---------------------------------------------------------------- template plumbing
def drop_slide(prs, index):
    sld_id = prs.slides._sldIdLst[index]
    prs.slides._sldIdLst.remove(sld_id)
    # Remove the relationship outright (drop_rel keeps it while anything else still names
    # the rId), so the old slide and its content are not written into the new file
    prs.part.rels.pop(sld_id.rId)


def drop_orphan_slides(prs):
    """The template file also carries slides that are not in its slide list; remove them."""
    listed = {s.rId for s in prs.slides._sldIdLst}
    for rId, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith("/slide") and rId not in listed:
            prs.part.rels.pop(rId)


def drop_sections(prs):
    """Remove the template's slide sections, which still list the deleted slides."""
    for ext in prs.part._element.findall(qn("p:extLst") + "/" + qn("p:ext")):
        if any(child.tag.endswith("sectionLst") for child in ext):
            ext.getparent().remove(ext)


def move_slide(prs, old, new):
    lst = prs.slides._sldIdLst
    el = lst[old]
    lst.remove(el)
    lst.insert(new, el)


def set_lines(shape, lines):
    """Replace the text of a shape line by line, keeping the formatting of each line's first run."""
    tf = shape.text_frame
    paras = list(tf.paragraphs)
    while len(paras) < len(lines):
        new = copy.deepcopy(paras[-1]._p)
        paras[-1]._p.addnext(new)
        paras = list(tf.paragraphs)
    for p, text in zip(paras, lines):
        runs = p.runs
        if runs:
            runs[0].text = text
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
        else:
            p.add_run().text = text
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)


# ---------------------------------------------------------------- building blocks
class Deck:
    def __init__(self, template):
        self.prs = Presentation(template)
        self.title_xml = copy.deepcopy(self.prs.slides[14].shapes[0]._element)
        self.blank = next(l for l in self.prs.slide_layouts if l.name == "BLANK")
        self.count = 0

    def slide(self, title, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        # The template's dropped slides can still be referenced from other parts, so give new
        # slides part names that cannot clash with theirs
        self.count += 1
        s.part.partname = PackURI(f"/ppt/slides/slide{900 + self.count}.xml")
        el = copy.deepcopy(self.title_xml)
        s.shapes._spTree.append(el)
        set_lines(s.shapes[-1], [title])
        if notes:
            s.notes_slide.notes_text_frame.text = notes
            s.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{900 + self.count}.xml")
        return s


def text(slide, x, y, w, h, paras, size=14, color=DARK, bold=False, align=None, bullet=False,
         anchor=MSO_ANCHOR.TOP, spacing=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        if align:
            p.alignment = align
        parts = para if isinstance(para, list) else [(para, bold)]
        if bullet:
            parts = [("•  ", False)] + list(parts)
        for t, b in parts:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = b
            r.font.color.rgb = color
    return tb


def heading(slide, x, y, w, t, size=15):
    return text(slide, x, y, w, 0.4, t, size=size, color=RED, bold=True)


def _border(cell, color="C9A227", width=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.find(qn(tag))
        if ln is None:
            ln = tcPr.makeelement(qn(tag), {})
            tcPr.append(ln)
        ln.set("w", str(width))
        for child in list(ln):
            ln.remove(child)
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr)
        ln.append(fill)


def table(slide, x, y, w, rows, widths=None, size=11, row_h=0.36, bold_first_col=False):
    n_rows, n_cols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(row_h * n_rows))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    widths = widths or [w / n_cols] * n_cols
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = HEAD_FILL if i == 0 else (ALT_FILL if i % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if i == 0 else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = RGBColor(0, 0, 0) if i == 0 else DARK
            r.font.bold = i == 0 or (bold_first_col and j == 0)
            _border(cell)
    return gf


def image(slide, path, x, y, w, h, align="center"):
    """Place an image inside the box (x, y, w, h), keeping its aspect ratio."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
    return pic


def caption(slide, x, y, w, t):
    return text(slide, x, y, w, 0.3, t, size=11, color=RGBColor(0x55, 0x55, 0x55), align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- content
def build(template):
    m = json.loads(config.METRICS_PATH.read_text())
    eda = json.loads((config.MODEL_DIR / "eda_summary.json").read_text())
    scale_path = config.MODEL_DIR / "scale_benchmark.json"
    scale = json.loads(scale_path.read_text()) if scale_path.exists() else []
    best = m["best_model"]
    bm = m["models"][best]
    bt = bm["test"]
    by = {k: {str(r["grp"]): r for r in v} for k, v in eda["by"].items()}

    d = Deck(template)
    prs = d.prs
    keep = {0, 1, len(prs.slides) - 1}
    for i in reversed(range(len(prs.slides))):
        if i not in keep:
            drop_slide(prs, i)
    drop_sections(prs)
    drop_orphan_slides(prs)

    # Title slide
    title = prs.slides[0]
    for sh in title.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("Department of Computer Science"):
            set_lines(sh, ["Department of Computer Science and Engineering",
                           "CA3: Project Submission and Presentation",
                           "Employee Attrition Prediction using HR Analytics",
                           "(Industry Domain: HR Analytics)",
                           "Session: 2026-27"])
        elif "Industry Mentor" in t:
            set_lines(sh, ["Technology:", "Apache Spark, MLlib, Streamlit"])
        elif "Institute Mentor" in t:
            set_lines(sh, ["         Project Guide: Dr. Priya Dasarwar"])

    outline = prs.slides[1]
    body = next(sh for sh in outline.shapes if sh.name == "Text Placeholder 2")
    set_lines(body, ["Introduction & Problem Statement", "Objectives & Scope",
                     "Literature Review", "Dataset & Exploratory Analysis (Spark SQL)",
                     "Methodology & System Architecture", "Spark MLlib Pipeline & Model Training",
                     "Results & Model Comparison", "Feature Importance & Scalability",
                     "Web Application Demo", "Business Insights, Conclusion & Future Scope",
                     "GitHub Repository & References"])

    # 3. Introduction
    s = d.slide("Introduction & Problem Statement",
                "Start with why attrition matters: cost of replacement, knowledge loss. HR usually reacts "
                "too late. Our system predicts risk early and explains it.")
    heading(s, 0.65, 0.95, 12, "Why employee attrition matters")
    text(s, 0.65, 1.35, 12.0, 1.6, [
        "Every resignation costs the organisation hiring, onboarding and training spend, months of lost "
        "productivity and the knowledge that leaves with the employee.",
        "HR usually learns that someone is unhappy only when the resignation arrives — too late to act.",
        "HR systems already hold the data needed to see the risk earlier: role, pay, overtime, travel, "
        "tenure, promotions and satisfaction surveys."], size=15, bullet=True)
    heading(s, 0.65, 3.25, 12, "Problem statement")
    text(s, 0.65, 3.65, 12.0, 1.2,
         "Given the HR records of an organisation's employees, predict which employees are likely to leave, "
         "identify the factors driving each person's risk, and present this to HR through an easy-to-use "
         "application — built on a scalable big-data pipeline (Apache Spark).", size=15)
    table(s, 0.65, 5.0, 12.0, [
        ["Type of problem", "Data", "Platform", "Output for HR"],
        ["Supervised binary classification (Left / Stayed), imbalanced classes",
         f"IBM HR Analytics — {eda['rows']:,} employees, 35 attributes, {eda['attrition_rate']}% left",
         "Apache Spark 3.5: PySpark, Spark SQL, MLlib", "Risk score, risk band, drivers and suggested actions"]],
        widths=[3.0, 3.2, 2.8, 3.0], size=12, row_h=0.5)

    # 4. Objectives
    s = d.slide("Objectives & Scope")
    rows = [["#", "Objective", "Delivered as"],
            ["1", "Analyse the factors associated with attrition using Apache Spark", "Spark SQL analysis of 12 factors + MLlib correlation matrix"],
            ["2", "Build a reusable Spark MLlib feature pipeline", "6-stage pipeline with in-pipeline feature engineering"],
            ["3", "Train, tune and compare four classifiers with class-imbalance handling", "LR, Decision Tree, Random Forest, GBT — 5-fold CV, class weights"],
            ["4", "Select model and threshold without touching the test set", "Stratified 85 / 15 split; CV + out-of-fold tuning; test used once"],
            ["5", "Explain every prediction in HR terms", "What-if risk drivers mapped to retention actions"],
            ["6", "Deploy a working application and prove scalability", "Streamlit web app; pipeline run on ~735k rows"]]
    table(s, 0.65, 1.05, 12.0, rows, widths=[0.5, 5.6, 5.9], size=13, row_h=0.55)
    heading(s, 0.65, 5.15, 12, "Scope")
    text(s, 0.65, 5.55, 12.0, 1.0,
         "End-to-end: data analysis → model development → evaluation → explanation → deployment → documentation. "
         "Uses a public, anonymised dataset; Spark runs in local mode, and the same code runs on a cluster unchanged.",
         size=14)

    # 5. Literature review
    s = d.slide("Literature Review",
                "Key message: the literature agrees on predictors and on tree ensembles, but rarely builds on a "
                "big-data engine, rarely keeps tuning separate from the test, and rarely delivers a usable app.")
    rows = [["Study", "Approach", "Key finding"],
            ["Saradhi & Palshikar (2011), Expert Syst. Appl.", "Employee churn as classification; NB, SVM, LR, DT, RF", "SVM best; models must be judged on the minority (leaver) class"],
            ["Alao & Adeyemo (2013)", "Decision trees: C4.5, C5.0, REPTree, CART", "Readable rules linking salary, length of service and age to attrition"],
            ["Punnoose & Ajit (2016), IJARAI", "Seven classifiers incl. XGBoost", "Gradient boosting best and robust to noisy HR data"],
            ["Zhao et al. (2019), IntelliSys / Springer", "Ten algorithms on small, medium, large HR data", "Boosted-tree ensembles most reliable across data sizes"],
            ["Fallucchi et al. (2020), Computers", "IBM HR dataset; six classifiers", "Overtime, income, total working years and age are key predictors"],
            ["Yahia et al. (2021), IEEE Access", "\"Big data to deep data\" people analytics; ensembles", "Ensembles with feature selection improve attrition prediction"],
            ["Zaharia et al. (2016); Meng et al. (2016)", "Apache Spark and Spark MLlib", "Unified in-memory engine; ML pipelines that scale out"]]
    table(s, 0.5, 1.0, 12.3, rows, widths=[3.4, 4.2, 4.7], size=11.5, row_h=0.5)
    heading(s, 0.5, 5.15, 12, "Research gap addressed by this project")
    text(s, 0.5, 5.55, 12.3, 1.0,
         "Most studies stop at a notebook score on one machine. This project (i) builds on a scalable big-data engine, "
         "(ii) keeps model / threshold selection separate from the final test, (iii) explains each prediction as HR "
         "actions, and (iv) delivers a working application.", size=13)

    # 6. Dataset
    s = d.slide("Dataset Description")
    rows = [["Group", "Attributes"],
            ["Target", "Attrition (Yes / No)"],
            ["Personal", "Age, Gender, MaritalStatus, DistanceFromHome, Education, EducationField"],
            ["Job", "Department, JobRole, JobLevel, JobInvolvement, BusinessTravel, OverTime, PerformanceRating, TrainingTimesLastYear"],
            ["Compensation", "MonthlyIncome, DailyRate, HourlyRate, MonthlyRate, PercentSalaryHike, StockOptionLevel"],
            ["Satisfaction", "Environment, Job, Relationship Satisfaction (1–4), WorkLifeBalance (1–4)"],
            ["Tenure", "TotalWorkingYears, NumCompaniesWorked, YearsAtCompany, YearsInCurrentRole, YearsSinceLastPromotion, YearsWithCurrManager"],
            ["Removed", "EmployeeCount, Over18, StandardHours (same value for every row); EmployeeNumber kept as ID only"]]
    table(s, 0.5, 1.0, 8.1, rows, widths=[1.6, 6.5], size=11, row_h=0.6)
    image(s, FIG / "01_class_balance.png", 8.8, 1.0, 4.2, 3.2)
    text(s, 8.8, 4.35, 4.2, 2.0, [
        [("Source: ", True), ("IBM HR Analytics Employee Attrition & Performance (Kaggle)", False)],
        [("Size: ", True), (f"{eda['rows']:,} employees × 35 attributes, no missing values", False)],
        [("Imbalance: ", True), (f"{eda['leavers']} left ({eda['attrition_rate']}%) — a model that always says "
                                 "\"stays\" is 84% accurate and useless", False)]], size=12)

    # 7. EDA I
    s = d.slide("Exploratory Analysis with Spark SQL (1/2)",
                "Explain the method: GROUP BY aggregations in Spark SQL. Red bars are above the company average.")
    text(s, 0.5, 0.9, 12.3, 0.5,
         "SELECT factor, COUNT(*), AVG(label) FROM hr GROUP BY factor  —  red bars are above the company-wide "
         f"attrition rate of {eda['attrition_rate']}% (dashed line)", size=12, color=RGBColor(0x55, 0x55, 0x55))
    image(s, FIG / "02_overtime.png", 0.4, 1.4, 4.6, 3.3)
    image(s, FIG / "03_job_role.png", 5.1, 1.35, 4.9, 3.6)
    image(s, FIG / "07_marital.png", 10.0, 1.4, 3.2, 2.6)
    ot, ms = by["OverTime"], by["MaritalStatus"]
    rr = eda["by"]["JobRole"]
    text(s, 0.5, 5.15, 12.3, 1.3, [
        [("Overtime: ", True), (f"{ot['Yes']['rate']}% of overtime workers leave vs {ot['No']['rate']}% of others — almost 3×", False)],
        [("Job role: ", True), (f"{rr[0]['grp']}s {rr[0]['rate']}% vs {rr[-1]['grp']}s {rr[-1]['rate']}%", False)],
        [("Marital status: ", True), (f"single {ms['Single']['rate']}% vs married {ms['Married']['rate']}%", False)]],
        size=13, bullet=True)

    # 8. EDA II
    s = d.slide("Exploratory Analysis with Spark SQL (2/2)")
    for i, (f, cap) in enumerate([("04_income.png", "Monthly income band"), ("05_tenure.png", "Years at company"),
                                  ("06_age.png", "Age band"), ("09_stock.png", "Stock option level")]):
        x = 0.35 + (i % 2) * 3.2
        y = 0.95 + (i // 2) * 2.75
        image(s, FIG / f, x, y, 3.1, 2.45)
    image(s, FIG / "11_correlation.png", 6.9, 0.95, 6.2, 4.6)
    text(s, 6.9, 5.6, 6.2, 1.0, "Correlation matrix computed with Spark MLlib. No single variable exceeds |0.17| "
         "with attrition, so a multivariate model is needed; income, level and experience are highly collinear.", size=11.5)
    text(s, 0.4, 6.45, 6.3, 0.5, "Highest risk: young, early-tenure, low-income, level-1 employees without stock options",
         size=12, bold=True, color=RED)

    # 9. Methodology
    s = d.slide("Methodology")
    image(s, FIG / "00_methodology.png", 0.4, 0.95, 12.5, 2.2)
    rows = [["Step", "What was done"],
            ["Data processing", "PySpark DataFrames; constant columns removed; label encoded; numerics cast to double"],
            ["Feature engineering", "IncomePerJobLevel, TenureRatio, YearsPerCompany, PromotionGap, SatisfactionIndex (SQLTransformer)"],
            ["Data split", f"Stratified 85 / 15, hash-based so it is identical on every run: {m['split']['development']['rows']} development · {m['split']['test']['rows']} test (used once)"],
            ["Class imbalance", "Leavers weighted by stayers / leavers ratio (≈5.2) through MLlib weightCol"],
            ["Tuning", "CrossValidator, 5 folds, ROC-AUC, ParamGridBuilder grid per model, parallelism 4"],
            ["Selection", "Best model by CV AUC; threshold that maximises F1 on out-of-fold predictions"],
            ["Final evaluation", "Test set used once: AUC, accuracy, precision, recall, F1, confusion matrix"]]
    table(s, 0.6, 3.3, 12.1, rows, widths=[2.6, 9.5], size=12, row_h=0.43)

    # 10. Architecture
    s = d.slide("System Architecture")
    image(s, FIG / "00_architecture.png", 0.6, 0.95, 12.1, 6.1)

    # 11. Pipeline
    s = d.slide("Spark MLlib Pipeline & Models")
    image(s, FIG / "00_ml_pipeline.png", 0.4, 0.95, 12.5, 2.25)
    rows = [["Model", "Hyper-parameter grid (5-fold CV)", "Best parameters found"]]
    grids = {"Logistic Regression": "regParam {0.001, 0.01, 0.1} × elasticNet {0, 0.5, 1}",
             "Decision Tree": "maxDepth {3, 5, 7} × minInstancesPerNode {5, 20}",
             "Random Forest": "numTrees {100, 200} × maxDepth {5, 8}",
             "Gradient-Boosted Trees": "maxIter {50, 100} × maxDepth {3, 5}, stepSize 0.1"}
    for name, r in m["models"].items():
        bp = ", ".join(f"{k}={v:g}" for k, v in r["best_params"].items())
        rows.append([name, grids[name], bp])
    table(s, 0.6, 3.35, 12.1, rows, widths=[2.9, 5.0, 4.2], size=12, row_h=0.45)
    text(s, 0.6, 5.75, 12.1, 0.9, "The whole pipeline (not only the classifier) is cross-validated, so encoders and the "
         "scaler are re-fitted inside every fold — no information leaks from the held-out folds. The saved "
         "PipelineModel accepts raw HR records directly.", size=13)

    # 12. Results table
    s = d.slide("Results — Model Comparison",
                f"Best model {best}. Explain AUC = probability a random leaver is ranked above a random stayer.")
    rows = [["Model", "CV AUC", "OOF AUC", "Test AUC", "Acc.", "Prec.", "Recall", "F1", "Thr."]]
    for name, r in m["models"].items():
        t = r["test"]
        rows.append([name + ("  ★" if name == best else ""), f"{r['cv_auc']:.3f}", f"{r['oof_auc']:.3f}",
                     f"{r['test_auc']:.3f}", f"{t['accuracy']:.3f}", f"{t['precision']:.3f}",
                     f"{t['recall']:.3f}", f"{t['f1']:.3f}", f"{t['threshold']:.2f}"])
    table(s, 0.5, 1.0, 7.3, rows, widths=[2.3] + [0.625] * 8, size=10.5, row_h=0.5, bold_first_col=True)
    image(s, FIG / "14_model_comparison.png", 7.95, 0.95, 5.1, 3.0)
    heading(s, 0.5, 3.75, 7, f"Selected model: {best}")
    text(s, 0.5, 4.15, 12.3, 2.4, [
        f"Highest cross-validated AUC ({bm['cv_auc']:.3f}) → chosen; test AUC {bm['test_auc']:.3f} on {sum(bt['confusion'].values())} unseen employees",
        f"Finds {100 * bt['recall']:.0f}% of the employees who actually left, with precision {100 * bt['precision']:.0f}% and F1 {100 * bt['f1']:.0f}% (threshold {bt['threshold']:.2f})",
        "Model and threshold were chosen on development data only — the test numbers are an unbiased estimate",
        "CV AUC = mean of 5 folds; OOF = out-of-fold predictions; test metrics at each model's tuned threshold"], size=13, bullet=True)

    # 13. ROC + CM
    s = d.slide("Results — ROC Curves & Confusion Matrix")
    image(s, FIG / "12_roc_curves.png", 0.4, 0.95, 4.3, 3.6)
    image(s, FIG / "13_pr_curves.png", 4.6, 0.95, 4.3, 3.6)
    image(s, FIG / "15_confusion_matrix.png", 9.0, 0.95, 4.0, 3.6)
    c = bt["confusion"]
    text(s, 0.6, 4.8, 12.1, 1.8, [
        f"Of {c['tp'] + c['fn']} test employees who left, the model correctly flagged {c['tp']} and missed {c['fn']}.",
        f"Of {c['tn'] + c['fp']} who stayed, it cleared {c['tn']} and raised {c['fp']} false alarms.",
        "Every model ranks employees far better than chance (diagonal). HR can move the threshold to trade "
        "precision for recall."], size=13, bullet=True)

    # 14. Feature importance
    s = d.slide("Key Drivers of Attrition")
    image(s, FIG / "16_feature_importance.png", 0.4, 0.95, 6.4, 5.4)
    heading(s, 7.0, 1.0, 6, "What the model learned")
    top, seen = [], set()
    for f in m["feature_importance"]:
        base = f["feature"].split(" = ")[0]
        if f["feature"] == "OverTime = No":
            continue  # mirror image of "OverTime = Yes"
        seen.add(base)
        arrow = "▲ raises risk" if f.get("signed", 0) > 0 else "▼ lowers risk"
        top.append(f"{f['feature']}  ({arrow})")
        if len(top) == 8:
            break
    text(s, 7.0, 1.45, 5.9, 2.6, [f"{i}. {t}" for i, t in enumerate(top, 1)], size=12.5, spacing=2)
    heading(s, 7.0, 4.15, 6, "Explaining one employee (what-if)")
    text(s, 7.0, 4.55, 5.9, 2.2,
         "For each factor, the app sets it to the typical value of employees who stayed and re-scores with Spark. "
         "The drop in leave probability = that factor's contribution, shown with a suggested HR action.", size=12.5)

    # 15. Scalability
    if scale:
        s = d.slide("Big Data — Scalability Test")
        rows = [["Rows", "Model", "Fit (s)", "Score (s)"]]
        for r in scale:
            rows.append([f"{r['rows']:,}", r["model"], f"{r['fit_seconds']:.1f}", f"{r['score_seconds']:.1f}"])
        table(s, 0.5, 1.0, 6.0, rows, widths=[1.5, 2.5, 1.0, 1.0], size=11, row_h=0.4)
        image(s, FIG / "17_scalability.png", 6.8, 1.0, 6.2, 3.8)
        text(s, 6.8, 5.0, 6.2, 1.8, [
            "Data enlarged up to 500× (≈735,000 rows) with ±5% noise on numeric columns",
            "Same Spark pipeline, no code change at any size",
            "Time grows far slower than data: Spark uses every core; on a cluster, every machine",
            "First (smallest) run includes Spark start-up and JVM warm-up"],
             size=12.5, bullet=True)

    # 16-17. Application
    if (SHOTS / "01_dashboard.png").exists():
        s = d.slide("Web Application — HR Dashboard & Prediction",
                    "Live demo: open the app (run_app.bat), show dashboard filters, then predict one high-risk employee.")
        image(s, SHOTS / "01_dashboard.png", 0.35, 0.95, 6.25, 5.6)
        image(s, SHOTS / "03_predict_result.png", 6.75, 0.95, 6.25, 5.6)
        caption(s, 0.35, 6.6, 6.25, "HR Dashboard: KPIs and filters (department, gender, overtime)")
        caption(s, 6.75, 6.6, 6.25, "Single prediction: risk gauge, drivers and suggested HR actions")
        s = d.slide("Web Application — Batch Scoring & Model Performance")
        image(s, SHOTS / "04_batch.png", 0.35, 0.95, 6.25, 5.6)
        image(s, SHOTS / "05_performance.png", 6.75, 0.95, 6.25, 5.6)
        caption(s, 0.35, 6.6, 6.25, "Batch prediction: upload CSV → risk bands → download results")
        caption(s, 6.75, 6.6, 6.25, "Model performance page: comparison, ROC, confusion matrix")

    # 18. Insights
    s = d.slide("Business Insights & Recommendations")
    tb = by["TenureBand"]
    rows = [["Finding (from Spark analysis)", "Recommended HR action"],
            [f"Overtime workers leave at {ot['Yes']['rate']}% vs {ot['No']['rate']}%", "Track overtime hours; rebalance workload or add headcount"],
            [f"First two years at the company: {tb['0-1']['rate']}% attrition", "Structured onboarding, buddy programme, 30/60/90-day check-ins"],
            [f"Income under 3k: {by['IncomeBand']['<3k']['rate']}%; job level 1: {by['JobLevel']['1.0']['rate']}%", "Benchmark entry-level pay; publish progression paths from level 1"],
            [f"No stock options: {by['StockOptionLevel']['0.0']['rate']}% attrition", "Stock options or retention bonus for high performers"],
            [f"Frequent travellers: {by['BusinessTravel']['Travel_Frequently']['rate']}%", "Rotate travel duties; travel allowance or remote options"],
            [f"Low job satisfaction: {by['JobSatisfaction']['1.0']['rate']}%; bad work-life balance: {by['WorkLifeBalance']['1.0']['rate']}%", "Pulse surveys, stay interviews, flexible hours"],
            [f"{rr[0]['grp']}s: {rr[0]['rate']}%", "Targeted retention plan: incentives, coaching, career path"]]
    table(s, 0.5, 1.0, 12.3, rows, widths=[5.6, 6.7], size=13, row_h=0.62)

    # 19. Conclusion
    s = d.slide("Conclusion & Future Scope")
    heading(s, 0.6, 0.95, 6, "Conclusion")
    text(s, 0.6, 1.4, 6.0, 5.0, [
        "Built a complete, working attrition prediction system on Apache Spark — from Spark SQL analysis to a deployed web app.",
        f"{best} selected: test ROC-AUC {bm['test_auc']:.3f}, recall {100 * bt['recall']:.0f}%, F1 {100 * bt['f1']:.0f}%.",
        "Overtime, income / level, tenure, stock options, travel and satisfaction drive attrition — most are actionable.",
        "Every prediction comes with its drivers and a suggested HR action.",
        "The same pipeline runs unchanged on ~735k rows — ready for real big-data volumes.",
        "All six objectives met."], size=13, bullet=True)
    heading(s, 6.9, 0.95, 6, "Future scope")
    text(s, 6.9, 1.4, 6.0, 5.0, [
        "Train on an organisation's own HR data with monthly snapshots (trends, not one snapshot)",
        "Join attendance, leave, survey and performance data in Spark",
        "Deploy on a Spark cluster (Databricks / EMR) with scheduled retraining and drift monitoring",
        "SHAP explanations to capture interactions between factors",
        "Fairness audit across gender and age before real use",
        "Survival analysis: predict when an employee may leave"], size=13, bullet=True)

    # 20. Repo
    s = d.slide("GitHub Repository & How to Run")
    heading(s, 0.6, 0.95, 12, "github.com/varunmundhada/employee-attrition-prediction")
    rows = [["Folder / file", "Contents"],
            ["src/", "config, Spark session, preprocessing, EDA, training, prediction, scalability benchmark"],
            ["app/streamlit_app.py", "Web application: dashboard, single and batch prediction, model performance"],
            ["models/", "Saved Spark PipelineModel, metrics.json, eda_summary.json"],
            ["reports/figures/", "All generated charts"],
            ["tests/", "pytest tests for cleaning, split and predictor"],
            ["report/ · ppt/", "LaTeX project report (SIT format) and this presentation"]]
    table(s, 0.6, 1.45, 7.2, rows, widths=[2.2, 5.0], size=11.5, row_h=0.5)
    text(s, 8.1, 1.45, 4.8, 4.5, [
        [("Run it", True)],
        "setup.bat — create venv, install packages",
        "run_pipeline.bat — EDA + training",
        "run_app.bat — open http://localhost:8501",
        "",
        [("Or manually", True)],
        "pip install -r requirements.txt",
        "python -m src.eda",
        "python -m src.train",
        "streamlit run app/streamlit_app.py"], size=12.5)

    # 21. References
    s = d.slide("References")
    refs = [
        "D. G. Allen, P. C. Bryant and J. M. Vardaman, \"Retaining talent: Replacing misconceptions with evidence-based strategies,\" Academy of Management Perspectives, vol. 24, no. 2, pp. 48–64, 2010.",
        "V. V. Saradhi and G. K. Palshikar, \"Employee churn prediction,\" Expert Systems with Applications, vol. 38, no. 3, pp. 1999–2006, 2011.",
        "D. Alao and A. B. Adeyemo, \"Analyzing employee attrition using decision tree algorithms,\" CISDIAR Journal, vol. 4, no. 1, pp. 17–28, 2013.",
        "R. Punnoose and P. Ajit, \"Prediction of employee turnover in organizations using machine learning algorithms,\" IJARAI, vol. 5, no. 9, 2016.",
        "Y. Zhao et al., \"Employee turnover prediction with machine learning: A reliable approach,\" IntelliSys 2018, AISC vol. 869, Springer, 2019.",
        "F. Fallucchi et al., \"Predicting employee attrition using machine learning techniques,\" Computers, vol. 9, no. 4, p. 86, 2020.",
        "N. B. Yahia, J. Hlel and R. Colomo-Palacios, \"From big data to deep data to support people analytics for employee attrition prediction,\" IEEE Access, vol. 9, 2021.",
        "M. Zaharia et al., \"Apache Spark: A unified engine for big data processing,\" Communications of the ACM, vol. 59, no. 11, 2016.",
        "X. Meng et al., \"MLlib: Machine learning in Apache Spark,\" JMLR, vol. 17, no. 34, 2016.",
        "IBM, \"IBM HR Analytics Employee Attrition & Performance,\" Kaggle dataset."]
    text(s, 0.6, 1.0, 12.2, 5.9, [f"{i}.  {r}" for i, r in enumerate(refs, 1)], size=13, spacing=7)

    # Thank-you slide (kept from the template) goes last
    move_slide(prs, 2, len(prs.slides) - 1)
    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print("Saved", OUT, f"({len(prs.slides)} slides)")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else TEMPLATE)
